#!/usr/bin/env python3
"""Очистка списка лидов: python parse_leads.py data.txt -o result.xlsx"""
import argparse
import re
from pathlib import Path

from openpyxl import Workbook

# Номер мобильного телефона РФ, с пробелами, скобками и дефисами.
PHONE = re.compile(r'(?<!\d)(?:(?:\+?7|8)[\s().-]*)?9\d{2}[\s().-]*\d{3}[\s().-]*\d{2}[\s().-]*\d{2}(?!\d)')
NAME = re.compile(r'[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?')
CHILD = re.compile(r'([А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?)\s+(\d{1,2})(?![\d:])')


def normalize_phone(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            return None
        value = str(int(value))
    digits = re.sub(r'\D', '', str(value))
    if len(digits) == 10 and digits.startswith('9'):
        return '7' + digits
    if len(digits) == 11 and digits[0] in '78' and digits[1] == '9':
        return '7' + digits[1:]
    return None


def read_text(path):
    for encoding in ('utf-8-sig', 'cp1251'):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError('Невозможно определить кодировку TXT')


def parse_txt(path):
    leads = {}  # телефон -> [родитель, список детей]
    issues = []
    parsed = 0
    for line_no, raw in enumerate(read_text(path).splitlines(), 1):
        line = ' '.join(raw.split())
        matches = list(PHONE.finditer(line))
        if not matches:
            continue  # дата, имя отправителя, сообщения без телефона
        if len(matches) != 1:
            issues.append((line_no, 'Несколько номеров в строке', raw))
            continue
        match = matches[0]
        phone = normalize_phone(match.group())
        if not phone:
            issues.append((line_no, 'Некорректный номер', raw))
            continue
        prefix = re.sub(r'^\d+\s*[.)]\s*', '', line[:match.start()].strip())
        parent_match = NAME.search(prefix)
        parent = parent_match.group() if parent_match else ''
        children = []
        for name, age in CHILD.findall(line[match.end():]):
            if 5 <= int(age) <= 17 and name.casefold() not in {c.casefold() for c in children}:
                children.append(name)
        if not parent or not children:
            issues.append((line_no, 'Нет имени родителя или ребёнка', raw))
        parsed += 1
        if phone not in leads:
            leads[phone] = [parent, children]
            continue
        saved_parent, saved_children = leads[phone]
        if saved_parent and parent and saved_parent.casefold() != parent.casefold():
            issues.append((line_no, 'Разные имена родителей у одного номера', raw))
        if not saved_parent:
            leads[phone][0] = parent
        known = {c.casefold() for c in saved_children}
        for child in children:
            if child.casefold() not in known:
                saved_children.append(child)
                known.add(child.casefold())
    return leads, issues, parsed


def write_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip('ABC', (22, 34, 80)):
        ws.column_dimensions[col].width = width
    if headers[0] == 'Телефон':
        for cell in ws['A'][1:]:
            cell.number_format = '@'


def main():
    parser = argparse.ArgumentParser(description='TXT с лидами -> Excel без дублей')
    parser.add_argument('source', type=Path, help='TXT с лидами')
    parser.add_argument('-o', '--output', type=Path, default=Path('result.xlsx'))
    args = parser.parse_args()
    if not args.source.is_file() or args.source.suffix.lower() != '.txt':
        parser.error('Укажи существующий .txt файл')
    if args.output.suffix.lower() != '.xlsx':
        parser.error('Результат должен быть .xlsx')
    if args.source.resolve() == args.output.resolve():
        parser.error('Выходной файл не должен совпадать с исходным')
    try:
        leads, issues, total = parse_txt(args.source)
        if not leads:
            raise ValueError('В TXT не найдено ни одного мобильного номера')
        wb = Workbook()
        wb.remove(wb.active)
        write_sheet(wb, 'Лиды', ['Телефон', 'Родитель', 'Ребёнок'],
                    ((p, v[0], ', '.join(v[1])) for p, v in leads.items()))
        write_sheet(wb, 'На проверку', ['Строка', 'Причина', 'Исходный текст'], issues)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.output)
        print(f'Записей: {total}; уникальных телефонов: {len(leads)}; на проверку: {len(issues)}')
        print(f'Сохранено: {args.output}')
    except (OSError, ValueError) as exc:
        parser.exit(1, f'Ошибка: {exc}\n')


if __name__ == '__main__':
    main()
