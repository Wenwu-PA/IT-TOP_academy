#!/usr/bin/env python3
"""Сверка готовых лидов с CRM: python compare_crm.py result.xlsx crm.xls -o compare.xlsx"""
import argparse
import csv
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from parse_leads import PHONE, normalize_phone, read_text, write_sheet

PHONE_COLUMN = re.compile(r'телефон|мобильн|контактный\s+номер|phone', re.I)


def phones_in_cell(value):
    if value is None or isinstance(value, bool):
        return set()
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            return set()
        value = str(int(value))
    text = str(value)
    result = {p for m in PHONE.finditer(text) if (p := normalize_phone(m.group()))}
    if not result:
        single = normalize_phone(text)
        if single:
            result.add(single)
    return result


def read_leads(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb['Лиды'] if 'Лиды' in wb.sheetnames else wb.active
        rows = iter(ws.values)
        header = [str(c or '').strip().casefold() for c in next(rows, ())]
        try:
            columns = [next(i for i, col in enumerate(header) if text in col)
                       for text in ('телефон', 'родител', 'реб')]
        except StopIteration as exc:
            raise ValueError('В базе лидов нужны столбцы: Телефон, Родитель, Ребёнок') from exc
        leads, issues = {}, []
        for row_no, row in enumerate(rows, 2):
            if not any(c is not None and str(c).strip() for c in row):
                continue
            values = [row[i] if i < len(row) else None for i in columns]
            numbers = phones_in_cell(values[0])
            if len(numbers) != 1:
                issues.append((row_no, 'Некорректный или неоднозначный телефон', str(row)))
                continue
            number = next(iter(numbers))
            parent = str(values[1] or '').strip()
            children = [n.strip() for n in str(values[2] or '').split(',') if n.strip()]
            if not parent or not children:
                issues.append((row_no, 'Не указаны родитель или ребёнок', str(row)))
            if number not in leads:
                leads[number] = [parent, children]
            else:
                old_parent, old_children = leads[number]
                if old_parent and parent and old_parent.casefold() != parent.casefold():
                    issues.append((row_no, 'Разные родители для одного номера', str(row)))
                if not old_parent:
                    leads[number][0] = parent
                known = {c.casefold() for c in old_children}
                for name in children:
                    if name.casefold() not in known:
                        old_children.append(name)
                        known.add(name.casefold())
        return leads, issues
    finally:
        wb.close()


def crm_sheets(path):
    suffix = path.suffix.lower()
    if suffix == '.xlsx':
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                yield ws.title, ws.values
        finally:
            wb.close()
    elif suffix == '.xls':
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError('Для .xls установи xlrd: pip install xlrd') from exc
        wb = xlrd.open_workbook(str(path), on_demand=True,
                                ignore_workbook_corruption=True, logfile=None)
        try:
            for ws in wb.sheets():
                yield ws.name, (ws.row_values(i) for i in range(ws.nrows))
        finally:
            wb.release_resources()
    elif suffix == '.csv':
        content = read_text(path)
        try:
            delimiter = csv.Sniffer().sniff(content[:4096], delimiters=';,\t').delimiter
        except csv.Error:
            delimiter = ';'
        yield path.name, csv.reader(content.splitlines(), delimiter=delimiter)
    else:
        raise ValueError('Формат CRM должен быть .xls, .xlsx или .csv')


def read_crm(path):
    numbers, rows_count, issues = set(), 0, []
    for sheet_name, contents in crm_sheets(path):
        rows = iter(contents)
        phone_cols = []
        for i, row in enumerate(rows):
            phone_cols = [j for j, cell in enumerate(row)
                          if PHONE_COLUMN.search(str(cell or ''))]
            if phone_cols or i >= 39:
                break
        if not phone_cols:
            issues.append(('', 'Не найдена колонка «Телефон»', sheet_name))
            continue
        for row_no, row in enumerate(rows, i + 2):
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            rows_count += 1
            for col in phone_cols:
                if col < len(row):
                    value = row[col]
                    found = phones_in_cell(value)
                    numbers.update(found)
                    if value is not None and str(value).strip() and not found:
                        issues.append((f'{sheet_name}:{row_no}', 'Не распознан телефон CRM', str(value)))
    if not numbers:
        raise ValueError('В CRM не найдены телефонные номера; сравнение отменено')
    return numbers, rows_count, issues


def main():
    parser = argparse.ArgumentParser(description='Поиск лидов, отсутствующих в CRM')
    parser.add_argument('leads', type=Path, help='Готовый Excel с лидами')
    parser.add_argument('crm', type=Path, help='Выгрузка CRM (.xls/.xlsx/.csv)')
    parser.add_argument('-o', '--output', type=Path, default=Path('compare.xlsx'))
    args = parser.parse_args()
    for source in (args.leads, args.crm):
        if not source.is_file():
            parser.error(f'Не найден файл: {source}')
    if args.leads.suffix.lower() != '.xlsx' or args.crm.suffix.lower() not in ('.xls', '.xlsx', '.csv'):
        parser.error('Лиды: .xlsx; CRM: .xls, .xlsx или .csv')
    if args.output.suffix.lower() != '.xlsx':
        parser.error('Результат должен быть .xlsx')
    if args.output.resolve() in (args.leads.resolve(), args.crm.resolve()):
        parser.error('Нельзя перезаписывать исходные файлы')
    try:
        leads, issues = read_leads(args.leads)
        if not leads:
            raise ValueError('Не найдено корректных лидов')
        crm, count, crm_issues = read_crm(args.crm)
        missing = [p for p in leads if p not in crm]
        present = [p for p in leads if p in crm]
        wb = Workbook()
        wb.remove(wb.active)

        def lead_rows(phones):
            return ((p, leads[p][0], ', '.join(leads[p][1])) for p in phones)

        header = ['Телефон', 'Родитель', 'Ребёнок']
        write_sheet(wb, 'Нет в CRM', header, lead_rows(missing))
        write_sheet(wb, 'Есть в CRM', header, lead_rows(present))
        write_sheet(wb, 'Сводка', ['Показатель', 'Значение'], [
            ('Уникальных лидов', len(leads)), ('Есть в CRM', len(present)),
            ('Нет в CRM', len(missing)), ('Номеров в CRM', len(crm)), ('Строк в CRM', count)])
        write_sheet(wb, 'На проверку', ['Строка', 'Причина', 'Исходный текст'], issues + crm_issues)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.output)
        print(f'Лидов: {len(leads)}; в CRM: {len(present)}; нет в CRM: {len(missing)}')
        print(f'Проверить: {len(issues) + len(crm_issues)}; сохранено: {args.output}')
    except (OSError, ValueError, RuntimeError) as exc:
        parser.exit(1, f'Ошибка: {exc}\n')


if __name__ == '__main__':
    main()
